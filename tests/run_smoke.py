# -*- coding: utf-8 -*-
"""Test dymny makra poza HyperMeshem (atrapa API + PyQt5 offscreen).

Uruchomienie (z katalogu repozytorium):
    QT_QPA_PLATFORM=offscreen python3 tests/run_smoke.py [folder_wynikow]

Sprawdza: diagnostykę, analizę metryk (MQ), deltę wielu metryk (widoki,
przełączanie, wyciszanie, wygaszanie), różne sposoby przezroczystości,
analizę powierzchni, raport porównawczy, automat (struktura folderów,
zrzuty, PPTX) oraz budowę okna: zrzut każdej karty i legendy do PNG.
"""
import io
import os
import shutil
import sys
import tempfile
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["HMQS_NO_GUI"] = "1"

import fake_hm                      # noqa: E402
import make_models                  # noqa: E402

OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(tempfile.gettempdir(), "hmqs_smoke"))
if os.path.isdir(OUT):
    shutil.rmtree(OUT)
os.makedirs(OUT)
os.environ["USERPROFILE"] = OUT      # ustawienia .hm_quality_studio.json do folderu testu
os.environ["HOME"] = OUT
fake_hm.install()

FAILS = []


def check(cond, msg):
    if cond:
        print("  OK   %s" % msg)
    else:
        print("  FAIL %s" % msg)
        FAILS.append(msg)


def section(t):
    print("\n=== %s ===" % t)


# --- makro jako modul (jak exec w HyperMeshu) ---
ns = {"__name__": "hmqs", "__file__": os.path.join(ROOT, "HM_Quality_Studio.py")}
src = io.open(ns["__file__"], encoding="ascii").read()
exec(compile(src, ns["__file__"], "exec"), ns)
G = lambda k: ns[k]                  # noqa: E731
HM, MQ, DELTA, SURF, REPORT, VIEWS, PRESENT, FLOW, MESH, BUS = [G(k) for k in ("HM", "MQ", "DELTA", "SURF", "REPORT", "VIEWS", "PRESENT", "FLOW", "MESH", "BUS")]
BUS.quiet = True
STATE = fake_hm.STATE

ref, inf = make_models.make_pair(os.path.join(OUT, "models"))

# --- 1. API + diagnostyka ---
section("API / diagnostyka")
check(HM.ok(), "API atrapy dostepne")
ok, msg = HM.read_file(inf)
check(ok, "read_file INF: %s" % msg)
check(HM.bulk_ok(), "odczyt hurtowy Tcl aktywny")
rows = HM.diagnostics()
bad = [r for r in rows if not r[1]]
for r in rows:
    print("     %s %s: %s" % ("+" if r[1] else "-", r[0], r[2]))
check(all(ok for name, ok, det in rows if "color_rgb" in name or "Przezroczysto" in name or "Transparency" in name), "diagnostyka: kolor RGB i przezroczystosc OK")
check(HM.rgb_exact(), "dokladne kolory RGB (color_rgb) wykryte")

# --- 2. MQ ---
section("Metryki i grupy (MQ)")
MQ.analyze()
check(len(MQ.analyzed) >= 3, "MQ analizowane metryki: %s" % MQ.analyzed)
n_before = len(STATE.elems)
MQ.apply_view("skew")
check(MQ.view == "skew", "MQ apply_view skew")
rows = MQ.legend_rows("skew")
check(all(r["rgb"] == STATE.comp_rgb(STATE.comp_by_name(r["comp"])) for r in rows if STATE.comp_by_name(r["comp"])), "legenda MQ = faktyczne kolory komponentow")
MQ.restore()
check(len(STATE.elems) == n_before and not MESH.active(), "MQ restore: elementy zachowane, stan czysty")
check(STATE.comp_by_name("MQ_SKEW_OK") == 0, "MQ restore: komponenty narzedzia usuniete")

# --- 3. Delta wielu metryk ---
section("Delta REF vs INF (wiele metryk)")
DELTA.ref_file, DELTA.inf_file = ref, inf
DELTA.metrics = {"ar": True, "jac": True, "skew": True, "disp": True}
DELTA.dim2, DELTA.dim3 = True, True
DELTA.set_mode("delta")
DELTA.fade_gray, DELTA.fade_level, DELTA.fade_style, DELTA.hide_others = True, 70, "white", True
DELTA.mark_extremes, DELTA.show_improved = True, True
DELTA.run()
check(DELTA.done and set(DELTA.analyzed) == {"ar", "jac", "skew", "disp"}, "delta policzona dla 4 metryk: %s" % DELTA.analyzed)
r = DELTA.res["ar"]
check(r.counts["band"] > 0 and r.counts["gray"] > 0 and r.counts["unm"] == 2, "AR: pasma %d, bez zmian %d, bez odpow. %d" % (r.counts["band"], r.counts["gray"], r.counts["unm"]))
check(DELTA.view == "ar", "pierwszy widok = AR")
gid = STATE.comp_by_name("D_dAR_bez_zmian")
check(gid and STATE.comps[gid]["transp"] == 70 and STATE.comp_rgb(gid) == (255, 255, 255), "bez zmian: biale + przezroczystosc 70%% (metoda %s)" % HM._transp_method)
check(DELTA.gray_faded == "transp", "stan wyciszenia: transp")
check(not STATE.comps[STATE.comp_by_name("beams")]["shown"], "komponent spoza narzedzia (beams) wygaszony")
lm = DELTA.legend_model("ar")
check(lm and len(lm["bands"]) == DELTA.band_count and lm["bands"][0]["rgb"] != lm["bands"][-1]["rgb"], "legenda AR: %d pasm, kolory rozne" % len(lm["bands"]))
first_cols = [b["rgb"] for b in lm["bands"]]
check(first_cols[0][2] > first_cols[0][0] and first_cols[-1][0] > first_cols[-1][2], "paleta ANSYS: dol niebieski, gora czerwona")
# przelaczenie widoku
DELTA.apply_view("jac")
check(DELTA.view == "jac", "apply_view jac")
ar_comps = [STATE.comp_by_name(n) for n in DELTA.res["ar"].comps]
check(all(c and not STATE.comps[c]["shown"] for c in ar_comps), "komponenty AR wygaszone po przelaczeniu na Jacobian")
check(all(len(STATE.elems_of_comps([c])) == 0 for c in ar_comps), "komponenty AR puste po przelaczeniu")
jac_gray = STATE.comp_by_name("D_dJac_bez_zmian")
check(jac_gray and STATE.comps[jac_gray]["transp"] == 70, "bez zmian Jacobian: przezroczystosc przeniesiona")
placed = sum(len(STATE.elems_of_comps([STATE.comp_by_name(n)])) for n in DELTA.res["jac"].comps if STATE.comp_by_name(n))
check(placed == len(DELTA.elems), "wszystkie elementy zakresu w komponentach Jacobian (%d)" % placed)
# przezroczystosc na zywo
DELTA.set_fade(True, 40, "gray")
check(STATE.comps[jac_gray]["transp"] == 40 and STATE.comp_rgb(jac_gray) == (163, 163, 163), "set_fade na zywo: 40%% + szare")
DELTA.set_fade(True, 90, "white")
check(STATE.comps[jac_gray]["transp"] == 90 and STATE.comp_rgb(jac_gray) == (255, 255, 255), "set_fade na zywo: 90%% + biale")
DELTA.set_hide_others(False)
check(STATE.comps[STATE.comp_by_name("beams")]["shown"], "set_hide_others(False): beams widoczne")
DELTA.set_hide_others(True)
txt = DELTA.inspect(5, "jac")
check("Q(REF)" in txt, "inspect po ID: %s" % txt.splitlines()[1][:60])
rows = DELTA.summary_rows()
check(len(rows) == 5, "summary_rows: 4 metryki")
# skala reczna + recompute
DELTA.manual_bounds["ar"] = [0.05, 0.5, 1.0, 2.0]
DELTA.auto_scale = False
DELTA.recompute("ar")
DELTA.apply_view("ar")
check(DELTA.res["ar"].k == 3 and DELTA.res["ar"].scale_mode == "manual", "reczna skala AR: 3 pasma")
DELTA.auto_scale = True
DELTA.manual_bounds = {}
DELTA.recompute()
# disp
rd = DELTA.res["disp"]
check(rd.max_info and rd.max_info[0] > 0, "przesuniecie: max %.3f mm (el. %s)" % (rd.max_info[0], rd.max_info[1]))
# restore
n_el = len(STATE.elems)
DELTA.restore()
check(len(STATE.elems) == n_el and not MESH.active() and STATE.comps[STATE.comp_by_name("beams")]["shown"], "restore: elementy zachowane, wygaszone wlaczone")
check(STATE.comp_by_name("D_dAR_bez_zmian") == 0, "restore: komponenty delty usuniete")

# --- 3b. przezroczystosc: inne wersje HM ---
section("Przezroczystosc: warianty HM")
for mode, read, expect in (("mark", False, "tclmark"), ("pyapi", False, "pyapi"), ("none", False, ""), ("setvalue", True, "setvalue")):
    fake_hm.CONFIG["transp_mode"], fake_hm.CONFIG["transp_read"] = mode, read
    HM._transp_method = ""
    cid = HM.ensure_comp("T_test")
    got = HM.set_transparency([cid], 55)
    check(got == expect, "tryb %s / odczyt %s -> metoda %r" % (mode, read, got))
    if expect:
        check(STATE.comps[cid]["transp"] == 55, "  poziom 55 ustawiony")
    HM.delete_comps([cid])
fake_hm.CONFIG["transp_mode"], fake_hm.CONFIG["transp_read"] = "none", False
HM._transp_method = ""
DELTA.run()
gid = STATE.comp_by_name("D_dAR_bez_zmian")
check(DELTA.gray_faded == "white" and STATE.comp_rgb(gid) == (255, 255, 255) and STATE.comps[gid]["shown"], "bez przezroczystosci w HM: bez zmian biale i widoczne")
DELTA.restore()
fake_hm.CONFIG["transp_mode"], fake_hm.CONFIG["transp_read"] = "setvalue", True
HM._transp_method = ""
# kolory bez color_rgb (starsze HM)
fake_hm.CONFIG["rgb"] = False
HM._rgb_ok = None
DELTA.run()
lm = DELTA.legend_model("ar")
gid = STATE.comp_by_name(DELTA.res["ar"].comps[0])
check(not HM.rgb_exact() and lm["bands"][-1]["rgb"] == STATE.comp_rgb(gid), "bez color_rgb: legenda = kolor z palety na komponencie")
DELTA.restore()
fake_hm.CONFIG["rgb"] = True
HM._rgb_ok = None

# --- 4. Powierzchnie ---
section("Powierzchnie - elementy krytyczne")
SURF.source, SURF.surf_text, SURF.dim3, SURF.dim2, SURF.min_shared = "surfs", "1", True, False, 3
SURF.models = "pair"
SURF.ref_file, SURF.inf_file = ref, inf
SURF.crit["aspectratio"] = {"use": True, "lo": None, "hi": 1.6}
SURF.crit["jacobian"] = {"use": True, "lo": 0.7, "hi": None}
SURF.crit["skew"] = {"use": True, "lo": None, "hi": 40.0}
SURF.crit["warpage"] = {"use": True, "lo": None, "hi": 1.5}
SURF.out_file = os.path.join(OUT, "surf", "powierzchnie.txt")
SURF.fmt = {"txt": True, "csv": True, "xlsx": True, "html": True}
msg = SURF.run()
print("     " + msg.splitlines()[0])
ids_top = SURF.results["INF"]["ids"]
check(len(ids_top) == 8 * 6, "elementy 3D na powierzchni 1 (gorna sciana): %d" % len(ids_top))
rows = SURF.summary_rows()
check(len(rows) == 5 and rows[1][3] != rows[1][5], "tabela REF / INF / delta (%s)" % " | ".join(str(c if not isinstance(c, dict) else c["t"]) for c in rows[1]))
check(all(os.path.isfile(f) for f in SURF.last_files) and len(SURF.last_files) == 4, "pliki powierzchni: %d" % len(SURF.last_files))
check(any(s["name"].startswith("SURF_INF_") for s in STATE.sets.values()), "zestawy poza tolerancja utworzone")
# inne zrodla
SURF.source, SURF.src_text = "comps", "skin_top"
SURF.models = "current"
SURF.run()
check(len(SURF.results["MODEL"]["ids"]) == 48, "zrodlo: komponent skin_top -> 48 elementow 3D")
SURF.source, SURF.src_text, SURF.dim2 = "elems", "1-10 12", True
SURF.run()
check(len(SURF.results["MODEL"]["ids"]) == 11, "zrodlo: ID elementow -> 11")
SURF.source, SURF.src_text = "sets", list(STATE.sets.values())[0]["name"]
SURF.run()
check(len(SURF.results["MODEL"]["ids"]) > 0, "zrodlo: zestaw")
fake_hm.CONFIG["by_geoms"] = False
SURF.source, SURF.surf_text = "surfs", "1"
try:
    SURF.run()
    check(False, "brak 'by geoms' powinien dac czytelny blad")
except ValueError as e:
    check("komponent" in "%s" % e or "component" in "%s" % e, "brak 'by geoms': czytelna podpowiedz")
fake_hm.CONFIG["by_geoms"] = True
check(G("parse_id_list")("12 13, 20-22; 13") == [12, 13, 20, 21, 22], "parse_id_list")

# --- 5. Automat ---
section("Automat REF vs INF")
from PyQt5 import QtWidgets, QtCore  # noqa: E402
app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])   # legendy SVG mierza tekst przez Qt
VIEWS.clear()
HM.read_file(inf)
VIEWS.remember()
VIEWS.views[-1]["name"] = "Widok_gora"
HM.set_view([0, 1, 0, 0, -1, 0, 0, 0, 0, 0, 1, 0, 0, 0, 0, 1, -1, -1, 1, 1])
VIEWS.remember()
VIEWS.views[-1]["name"] = "Widok_bok"
FLOW.ref_file, FLOW.inf_file, FLOW.out_dir, FLOW.name, FLOW.stamp = ref, inf, os.path.join(OUT, "wyniki"), "test_run", False
FLOW.do_delta = FLOW.do_report = FLOW.do_surf = FLOW.do_pptx = True
SURF.source, SURF.surf_text, SURF.dim2, SURF.dim3 = "surfs", "1", False, True
REPORT.fmt = {"txt": True, "csv": True, "xlsx": True, "html": True}
REPORT.use_metric["warpage"] = True
DELTA.metrics = {"ar": True, "jac": True, "skew": True, "disp": False}
PRESENT.preview_on = False
info = FLOW.run(deliver=True)
root = info["root"]
print("     " + "\n     ".join(info["log"]))
for sub in ("01_delta", "02_raport_jakosci", "03_powierzchnie", "04_prezentacja"):
    check(os.path.isdir(os.path.join(root, sub)), "podfolder %s" % sub)
shots = [p for p in info["files"] if p.endswith(".png") and "01_delta" in p]
check(len(shots) == 3 * 2, "zrzuty delty: %d (3 metryki x 2 widoki)" % len(shots))
check(all(os.path.isfile(os.path.join(root, "01_delta", t, "legenda_%s.svg" % t)) for t in ("dAR", "dJac", "dSkew")), "legendy SVG per metryka")
check(os.path.isfile(os.path.join(root, "podsumowanie.txt")) and os.path.isfile(os.path.join(root, "index.html")), "podsumowanie.txt + index.html")
pptx_path = os.path.join(root, "04_prezentacja", "test_run.pptx")
check(os.path.isfile(pptx_path), "prezentacja zapisana")
import pptx as _pptx                # noqa: E402
prs = _pptx.Presentation(pptx_path)
check(len(prs.slides) == 1 + 6 + 1 + 1 + 1, "slajdy: %d (tytul + 6 delta + zbiorczy + raport + powierzchnie)" % len(prs.slides))
check(any(f.endswith("_porownanie.xlsx") or f.endswith("_comparison.xlsx") for f in info["files"]), "raport porownawczy XLSX")
check(DELTA.done and DELTA.view == "ar" and MESH.owner == "delta", "po automacie: delta AR na siatce")
from PIL import Image               # noqa: E402
im = Image.open(shots[0])
check(im.size[0] > 100, "zrzut delty to obraz %dx%d" % im.size)
# ustawienia
check(G("save_settings")() and G("load_settings")(), "zapis / odczyt ustawien")
DELTA.restore()

# --- 6. GUI ---
section("GUI (offscreen)")
ns["yes_no"] = lambda *a, **k: True
ns["msg_box"] = lambda *a, **k: None
win = None
try:
    win = G("StudioWindow")()
    PRESENT.open_cb = lambda m, p: None        # bez modalnej galerii po eksporcie
    win.resize(1380, 900)
    win.show()
    app.processEvents()
    shots_dir = os.path.join(OUT, "gui")
    os.makedirs(shots_dir, exist_ok=True)
    for i, key in enumerate(win.TAB_KEYS):
        win.goto(key)
        app.processEvents()
        p = os.path.join(shots_dir, "%02d_%s.png" % (i, key))
        win.grab().save(p)
        check(os.path.getsize(p) > 1000, "zrzut karty %s" % key)
    # delta przez GUI (runner) + legenda
    win.goto("delta")
    win.tab_delta.load_from_engine()
    win.tab_delta.run()
    app.processEvents()
    check(DELTA.done, "delta uruchomiona z karty")
    win.tab_delta.show_view("skew")
    app.processEvents()
    check(DELTA.view == "skew", "przelaczenie metryki z karty")
    win.tab_delta.sp_fade.setValue(33)
    app.processEvents()
    gid = STATE.comp_by_name("D_dSkew_bez_zmian")
    check(STATE.comps[gid]["transp"] == 33, "suwak przezroczystosci dziala na zywo (33%)")
    win.show_legend("delta")
    app.processEvents()
    leg = win.legends["delta"]
    leg.grab().save(os.path.join(shots_dir, "legend_delta.png"))
    for th in ("white", "none"):
        leg.set_theme(th)
        app.processEvents()
        leg.grab().save(os.path.join(shots_dir, "legend_delta_%s.png" % th))
    check(leg.isVisible(), "okno legendy delty")
    svg = G("delta_legend_prims")("dark", "skew").svg("#41566b")
    check("<svg" in svg and svg.count("<rect") > 10, "legenda SVG")
    win.goto("delta")
    app.processEvents()
    win.grab().save(os.path.join(shots_dir, "03_delta_after.png"))
    win.goto("start")
    app.processEvents()
    win.grab().save(os.path.join(shots_dir, "00_start_after.png"))
    # surf przez GUI
    win.goto("surf")
    win.tab_surf.rb_cur.setChecked(True)
    win.tab_surf.rb_src["surfs"].setChecked(True)
    win.tab_surf.e_surf.setText("1")
    win.tab_surf.run()
    app.processEvents()
    check(SURF.last is not None, "powierzchnie z karty")
    win.grab().save(os.path.join(shots_dir, "04_surf_after.png"))
    # edytor skali delty
    ed = G("DeltaScaleEditor")(win, "skew")
    ed.show()
    app.processEvents()
    ed.grab().save(os.path.join(shots_dir, "scale_editor.png"))
    ed.close()
    # podglad slajdu delty (render)
    items = PRESENT.delta_all_items() + [PRESENT.delta_summary_item(), PRESENT.surface_item()]
    pngs = G("render_items_png")(items, 900)
    check(len(pngs) == len(items), "render slajdow: %d" % len(pngs))
    for i, p in enumerate(pngs[:3]):
        shutil.copy2(p, os.path.join(shots_dir, "slide_%d.png" % i))
    G("Presenter").drop_items(items)
    # automat przez GUI (bez dialogow)
    win.goto("flow")
    win.tab_flow.e_name.setText("gui_run")
    win.tab_flow.cb_prev.setChecked(False)
    win.tab_flow.cb_open.setChecked(False)
    win.tab_flow.run()
    app.processEvents()
    check(FLOW.last and FLOW.last["name"] == "gui_run" and os.path.isfile(FLOW.last["pptx"]), "automat z karty: %s" % (FLOW.last or {}).get("root"))
    win.grab().save(os.path.join(shots_dir, "01_flow_after.png"))
    win.switch_lang("en")
    app.processEvents()
    win2 = ns["builtins"]._HMQS_STUDIO
    win2.goto("delta")
    app.processEvents()
    win2.grab().save(os.path.join(shots_dir, "03_delta_en.png"))
    win2.close()
except Exception:
    traceback.print_exc()
    FAILS.append("GUI exception")
finally:
    try:
        if win is not None:
            win.close()
    except Exception:
        pass

print("\n%s" % ("ALL OK" if not FAILS else "FAILURES: %d\n  - %s" % (len(FAILS), "\n  - ".join(FAILS))))
print("wyniki: %s" % OUT)
sys.exit(1 if FAILS else 0)
